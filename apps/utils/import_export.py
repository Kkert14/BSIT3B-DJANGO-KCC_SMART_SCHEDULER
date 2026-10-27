"""Excel/CSV export and import utilities for the Smart Scheduler apps.
Requires openpyxl (already in requirements.txt).

The export and import sides share one canonical header format so that a file
produced by ``queryset_to_excel`` can always be read back by
``excel_to_dict_list``.

Canonical header rules
----------------------
A header cell is normalised to lowercase with non-alphanumeric characters
collapsed to a single underscore, and Django's ``__`` relationship separator
collapsed to a single underscore. So all of these refer to the same column::

    "Teacher Employee Number"   -> teacher_employee_number
    "teacher__employee_number"   -> teacher_employee_number
    "teacher employee number"    -> teacher_employee_number
"""

import csv
import io
import re
from datetime import datetime, time

from django.http import HttpResponse
from django.utils import timezone as django_timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Matches any run of characters that should not appear in a header key.
_NON_KEY_CHARS = re.compile(r'[^a-z0-9]+')

# Formats accepted when parsing a time cell.
_TIME_FORMATS = ('%H:%M', '%H:%M:%S')

# Values that are interpreted as True when parsing a boolean cell.
_TRUTHY = ('true', '1', 'yes', 'y', 't')


def normalize_header(value):
    """Return the canonical key for a header cell.

    >>> normalize_header('Teacher Employee Number')
    'teacher_employee_number'
    >>> normalize_header('teacher__employee_number')
    'teacher_employee_number'
    """
    text = str(value or '').strip().lower()
    text = text.replace('__', ' ')
    text = _NON_KEY_CHARS.sub('_', text)
    return text.strip('_')


def get_value(row, field, default=''):
    """Return the value of ``field`` in a parsed row.

    ``field`` may be written with Django's relationship separator
    (``'teacher__employee_number'``) or as the canonical header form
    (``'teacher_employee_number'``); both resolve to the same key.
    """
    key = normalize_header(field)
    if key in row:
        value = row[key]
        return default if value is None else value
    return default


def get_str(row, field, default=''):
    """Return the value of ``field`` as a stripped string.

    Blank cells come back as ``''`` rather than the string ``'None'``.
    """
    value = get_value(row, field, default)
    if value is None:
        return default
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def parse_bool(value, default=True):
    """Interpret a cell value as a boolean."""
    if value is None or value == '':
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in _TRUTHY


def parse_time(value):
    """Parse a cell value into a ``datetime.time``.

    Accepts a ``time``, a ``datetime``, or a string in ``HH:MM`` or
    ``HH:MM:SS`` form. Returns ``None`` when the value cannot be parsed.
    """
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in _TIME_FORMATS:
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return None


def queryset_to_excel(queryset, field_names, sheet_name="Export", title=None):
    """
    Convert a Django queryset to an Excel workbook.

    Args:
        queryset: Django QuerySet to export
        field_names: List of field names (supports __ for related fields, e.g. 'building__code')
        sheet_name: Name of the worksheet
        title: Optional title for the first row

    Returns:
        openpyxl Workbook object
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Add title row
    row = 1
    if title:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(field_names))
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(bold=True, size=14)
        cell.alignment = Alignment(horizontal='center')
        row = 2

    # Add header row
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, field_name in enumerate(field_names, start=1):
        cell = ws.cell(row=row, column=col_idx, value=field_name.replace('__', ' ').title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # Add data rows
    for obj in queryset:
        row += 1
        for col_idx, field_name in enumerate(field_names, start=1):
            # Support related field access via __ (e.g. 'building__code')
            value = obj
            for part in field_name.split('__'):
                value = getattr(value, part, None)
                if value is None:
                    break
            cell = ws.cell(row=row, column=col_idx, value=value if value is not None else '')
            # Convert timezone-aware datetimes to naive for openpyxl compatibility
            if isinstance(value, datetime) and value.tzinfo is not None:
                cell.value = django_timezone.make_naive(value)
            cell.alignment = Alignment(vertical='top')

    # Auto-size columns
    for col_idx in range(1, len(field_names) + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        for row_idx in range(1, row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

    return wb


def workbook_to_response(wb, filename):
    """
    Convert an openpyxl Workbook to an HttpResponse for download.
    """
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def csv_to_dict_list(file_content):
    """
    Parse CSV file content into a list of dictionaries.
    """
    text = file_content.decode('utf-8')
    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def _is_title_row(row, next_row):
    """True when ``row`` looks like a merged title rather than a header row.

    A title row has a single filled cell (the title spans the sheet) and the
    following row is the real header row with more than one filled cell.
    """
    filled = [v for v in row if v is not None and str(v).strip() != '']
    if len(filled) != 1:
        return False
    next_filled = [v for v in (next_row or []) if v is not None and str(v).strip() != '']
    return len(next_filled) > 1


def excel_to_dict_list(file_content):
    """
    Parse Excel file content into a list of dictionaries.

    Header cells are normalised to canonical keys (see ``normalize_header``),
    so a file written by ``queryset_to_excel`` round-trips cleanly. A leading
    title row is detected and skipped.
    """
    from openpyxl import load_workbook
    buffer = io.BytesIO(file_content)
    wb = load_workbook(buffer)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_index = 0
    if len(rows) > 1 and _is_title_row(rows[0], rows[1]):
        header_index = 1

    headers = [normalize_header(h) for h in rows[header_index]]
    if not any(headers):
        return []

    result = []
    for row in rows[header_index + 1:]:
        if not any(row):
            continue
        result.append(dict(zip(headers, row)))

    return result
