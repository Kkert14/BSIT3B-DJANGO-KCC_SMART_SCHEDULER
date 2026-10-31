"""Tests for the schedule report pages and Excel export."""

import json
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.messages.middleware import MessageMiddleware
from django.contrib.sessions.middleware import SessionMiddleware
from django.test import RequestFactory, TestCase
from django.urls import resolve, reverse
from openpyxl import load_workbook

from apps.building.models import Building
from apps.curriculum.models import Curriculum, Semester
from apps.faculty.models import Faculty
from apps.program.models import Program
from apps.rooms.models import Room
from apps.schedule.models import Schedule
from apps.section.models import Section
from apps.subject.models import Subject
from apps.users.models import UserProfile

User = get_user_model()


class ReportTestMixin:
    @classmethod
    def make_user(cls, username, role):
        user = User.objects.create_user(username=username, password='test-pass-123')
        profile, _ = UserProfile.objects.get_or_create(
            user=user, defaults={'role': role}
        )
        profile.role = role
        profile.save()
        user._state.fields_cache.pop('profile', None)
        return user

    def build_request(self, user, url_name, data=None):
        url = reverse(url_name)
        request = RequestFactory().get(url, data or {})
        SessionMiddleware(lambda r: None).process_request(request)
        request.session.save()
        MessageMiddleware(lambda r: None).process_request(request)
        request.user = user
        return request

    def call(self, user, url_name, data=None):
        request = self.build_request(user, url_name, data)
        match = resolve(request.path)
        return match.func(request, **match.kwargs)


class ReportTests(ReportTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.curriculum = Curriculum.objects.create(
            code='2025-2026', name='Curriculum 2025-2026',
            academic_year='2025-2026', is_active=True,
        )
        cls.semester = Semester.objects.create(
            curriculum=cls.curriculum, semester_number=1, is_active=True,
        )
        cls.program = Program.objects.create(
            program_code='BSIT', program_description='BS Information Technology',
        )
        cls.section = Section.objects.create(
            name='BSIT 1-A', program=cls.program, semester=cls.semester,
            year_level=1, capacity=40,
        )
        cls.section_two = Section.objects.create(
            name='BSIT 1-B', program=cls.program, semester=cls.semester,
            year_level=1, capacity=40,
        )
        cls.building = Building.objects.create(code='B1', name='Main Building')
        cls.room = Room.objects.create(
            building=cls.building, room_name='101', capacity=45,
        )
        cls.subject = Subject.objects.create(
            code='IT101', name='Introduction to Computing', units=3,
            lecture_hours=3, program=cls.program, semester=cls.semester,
            year_level=1,
        )
        cls.subject_two = Subject.objects.create(
            code='IT102', name='Computer Programming', units=3,
            lab_hours=1, program=cls.program, semester=cls.semester,
            year_level=1,
        )
        cls.teacher_user = cls.make_user('report_teacher', 'teacher')
        cls.other_teacher_user = cls.make_user('report_other', 'teacher')
        cls.unlinked_teacher = cls.make_user('report_unlinked', 'teacher')
        cls.admin = cls.make_user('report_admin', 'admin')
        cls.student = cls.make_user('report_student', 'student')

        cls.teacher = Faculty.objects.create(
            user=cls.teacher_user, employee_number='F-001',
            first_name='Maria', last_name='Santos', position='Instructor',
        )
        cls.other_teacher = Faculty.objects.create(
            user=cls.other_teacher_user, employee_number='F-002',
            first_name='Jose', last_name='Reyes', position='Instructor',
        )
        cls.schedule = Schedule.objects.create(
            semester=cls.semester, teacher=cls.teacher, subject=cls.subject,
            section=cls.section, room=cls.room, class_type='Lecture',
            day_of_week=1, start_time='08:00', end_time='09:30',
        )
        cls.other_schedule = Schedule.objects.create(
            semester=cls.semester, teacher=cls.other_teacher,
            subject=cls.subject_two, section=cls.section_two, room=cls.room,
            class_type='Laboratory', day_of_week=3,
            start_time='09:00', end_time='10:00',
        )

    def test_student_is_redirected_away(self):
        for url_name in ('schedule:schedule_reports', 'schedule:schedule_reports_export'):
            with self.subTest(url=url_name):
                response = self.call(self.student, url_name)
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response['Location'], reverse('dashboard'))

    def test_teacher_without_faculty_link_is_redirected(self):
        response = self.call(self.unlinked_teacher, 'schedule:schedule_reports')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response['Location'], reverse('schedule:schedule_list'))

    def test_master_report_lists_every_schedule(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports', {'report': 'master'}
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Master list', body)
        self.assertIn('IT101', body)
        self.assertIn('IT102', body)
        self.assertIn('Santos', body)
        self.assertIn('Reyes', body)

    def test_timetable_report_defaults_to_a_section(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports', {'report': 'timetable'}
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Section timetable', body)
        self.assertIn('BSIT 1-A', body)
        self.assertIn('IT101', body)
        self.assertNotIn('IT102', body)

    def test_grid_view_renders_a_grid(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports',
            {'report': 'timetable', 'view': 'grid'},
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('report-grid', body)
        self.assertIn('report-event', body)

    def test_teacher_only_sees_their_own_schedule(self):
        response = self.call(
            self.teacher_user, 'schedule:schedule_reports',
            {'report': 'master'},
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('IT101', body)
        self.assertNotIn('IT102', body)
        self.assertNotIn('<h2>Master list</h2>', body)

    def test_faculty_report_filters_a_single_teacher(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports',
            {'report': 'faculty', 'teacher': str(self.other_teacher.id)},
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Jose Reyes', body)
        self.assertNotIn('IT101', body)
        self.assertNotIn('Faculty load summary', body)

    def test_all_faculty_report_includes_load_summary(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports', {'report': 'faculty'}
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Faculty load summary', body)
        self.assertIn('Maria Santos', body)
        self.assertIn('Jose Reyes', body)

    def test_export_returns_a_workbook(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports_export', {'report': 'master'}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=5, column=1).value, 'Day')
        values = [
            sheet.cell(row=row, column=5).value for row in range(6, sheet.max_row + 1)
        ]
        self.assertIn('IT101', values)
        self.assertIn('IT102', values)

    def test_export_for_teacher_only_contains_own_classes(self):
        response = self.call(
            self.teacher_user, 'schedule:schedule_reports_export'
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        values = [
            sheet.cell(row=row, column=5).value for row in range(6, sheet.max_row + 1)
        ]
        self.assertIn('IT101', values)
        self.assertNotIn('IT102', values)


class ActiveTermTests(ReportTestMixin, TestCase):
    """Scheduling and reports follow the navbar term, not URL parameters."""

    @classmethod
    def setUpTestData(cls):
        cls.old_curriculum = Curriculum.objects.create(
            code='2024-2025', name='Curriculum 2024-2025',
            academic_year='2024-2025', is_active=False,
        )
        cls.old_semester = Semester.objects.create(
            curriculum=cls.old_curriculum, semester_number=1, is_active=False,
        )
        cls.curriculum = Curriculum.objects.create(
            code='2025-2026', name='Curriculum 2025-2026',
            academic_year='2025-2026', is_active=True,
        )
        cls.semester = Semester.objects.create(
            curriculum=cls.curriculum, semester_number=1, is_active=True,
        )
        cls.program = Program.objects.create(
            program_code='BSIT', program_description='BS Information Technology',
        )
        cls.old_section = Section.objects.create(
            name='BSIT 2-A', program=cls.program, semester=cls.old_semester,
            year_level=2, capacity=40,
        )
        cls.section = Section.objects.create(
            name='BSIT 1-A', program=cls.program, semester=cls.semester,
            year_level=1, capacity=40,
        )
        cls.old_subject = Subject.objects.create(
            code='OLD101', name='Previous Term Course', units=3,
            lecture_hours=3, program=cls.program, semester=cls.old_semester,
            year_level=2,
        )
        cls.subject = Subject.objects.create(
            code='IT101', name='Introduction to Computing', units=3,
            lecture_hours=3, program=cls.program, semester=cls.semester,
            year_level=1,
        )
        cls.building = Building.objects.create(code='B1', name='Main Building')
        cls.room = Room.objects.create(
            building=cls.building, room_name='101', capacity=45,
        )
        cls.admin = cls.make_user('term_admin', 'admin')
        cls.teacher = Faculty.objects.create(
            employee_number='F-010', first_name='Ana', last_name='Cruz',
            position='Instructor',
        )
        Schedule.objects.create(
            semester=cls.old_semester, teacher=cls.teacher, subject=cls.old_subject,
            section=cls.old_section, room=cls.room, class_type='Lecture',
            day_of_week=2, start_time='08:00', end_time='09:00',
        )
        cls.schedule = Schedule.objects.create(
            semester=cls.semester, teacher=cls.teacher, subject=cls.subject,
            section=cls.section, room=cls.room, class_type='Lecture',
            day_of_week=1, start_time='08:00', end_time='09:00',
        )

    def test_schedule_page_is_scoped_to_the_active_term(self):
        response = self.call(self.admin, 'schedule:schedule_list')
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('ss-grid-filters', body)
        self.assertIn('A.Y. 2025-2026 · 1st Sem', body)
        self.assertIn('BSIT 1-A', body)
        self.assertNotIn('BSIT 2-A', body)
        self.assertNotIn('OLD101', body)

    def test_schedule_page_drops_term_pickers(self):
        response = self.call(self.admin, 'schedule:schedule_list')
        body = response.content.decode()
        self.assertNotIn('id="f_curriculum"', body)
        self.assertNotIn('id="f_semester"', body)
        self.assertNotIn('name="year"', body)
        self.assertNotIn('name="semester"', body)

    def test_reports_ignore_year_and_semester_parameters(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports',
            {
                'report': 'master',
                'year': str(self.old_curriculum.id),
                'semester': '2',
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('IT101', body)
        self.assertNotIn('OLD101', body)
        self.assertIn('A.Y. 2025-2026 · 1st Sem', body)

    def test_batch_save_uses_the_active_term(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('schedule:schedule_batch_save_ajax'),
            {
                'teacher': str(self.teacher.id),
                'subject': str(self.subject.id),
                'section': str(self.section.id),
                'room': str(self.room.id),
                'class_type': 'Lecture',
                'semester': str(self.old_semester.id),
                'rows': json.dumps([
                    {'day': 5, 'start': '13:00', 'end': '14:00'},
                ]),
            },
        )
        self.assertEqual(response.status_code, 200, response.content)
        saved = Schedule.objects.get(
            subject=self.subject, day_of_week=5, start_time='13:00'
        )
        self.assertEqual(saved.semester_id, self.semester.id)

    def test_batch_save_rejects_a_subject_from_another_term(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('schedule:schedule_batch_save_ajax'),
            {
                'teacher': str(self.teacher.id),
                'subject': str(self.old_subject.id),
                'section': str(self.section.id),
                'room': str(self.room.id),
                'class_type': 'Lecture',
                'rows': json.dumps([
                    {'day': 5, 'start': '15:00', 'end': '16:00'},
                ]),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('not offered in the active term', response.json()['error'])


class NoActiveTermTests(ReportTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.curriculum = Curriculum.objects.create(
            code='2025-2026', name='Curriculum 2025-2026',
            academic_year='2025-2026', is_active=False,
        )
        cls.semester = Semester.objects.create(
            curriculum=cls.curriculum, semester_number=1, is_active=False,
        )
        cls.admin = cls.make_user('no_term_admin', 'admin')

    def test_schedule_page_shows_a_prompt_instead_of_filters(self):
        response = self.call(self.admin, 'schedule:schedule_list')
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('No active term selected', body)
        self.assertNotIn('id="scheduleForm"', body)

    def test_reports_show_a_prompt(self):
        response = self.call(
            self.admin, 'schedule:schedule_reports', {'report': 'master'}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('No active term selected', response.content.decode())
